from __future__ import annotations

import io
from collections.abc import Iterable, Iterator

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.shapes.base import BaseShape

from ..builder import ElementBuilder, table_row_text
from ..errors import ParserError


def parse_pptx(data: bytes, builder: ElementBuilder, max_slides: int) -> None:
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ParserError(
            "unreadable_pptx",
            f"PPTX 解析失败：{type(exc).__name__}。文件可能损坏或不是有效的 PPTX。",
            422,
        ) from exc

    slides = list(presentation.slides)
    if not slides:
        raise ParserError("empty_pptx", "PPTX 不包含任何幻灯片。", 422)
    if len(slides) > max_slides:
        builder.warn(
            f"PPTX 共 {len(slides)} 页，超过单次解析上限 {max_slides} 页，"
            f"仅解析前 {max_slides} 页。"
        )

    chart_warned = False
    notes_warned = False
    table_group = 0

    for slide_number, slide in enumerate(slides[:max_slides], start=1):
        builder.reset_heading_path()
        slide_start = len(builder.elements)
        shapes = list(_iter_shapes(slide.shapes))
        title_shape = _find_title_shape(shapes)
        title_shape_id = id(title_shape) if title_shape is not None else None
        paragraph_index = 0

        if title_shape is not None and title_shape.has_text_frame:
            title_paragraphs = [
                paragraph
                for paragraph in title_shape.text_frame.paragraphs
                if paragraph.text.strip()
            ]
            if title_paragraphs:
                paragraph_index += 1
                builder.add_heading(
                    title_paragraphs[0].text,
                    1,
                    page=slide_number,
                    paragraph=paragraph_index,
                )
                for paragraph in title_paragraphs[1:]:
                    paragraph_index += 1
                    builder.add(
                        "paragraph",
                        paragraph.text,
                        page=slide_number,
                        paragraph=paragraph_index,
                        level=(paragraph.level + 1) if paragraph.level else None,
                    )

        for shape in shapes:
            if shape.has_chart:
                if not chart_warned:
                    builder.warn("PPTX 图表数据未提取，仅跳过图表对象。")
                    chart_warned = True
                continue

            if shape.has_table:
                table_group += 1
                _emit_table(
                    shape.table, builder, page=slide_number, table_group=table_group
                )
                continue

            if not shape.has_text_frame:
                continue
            if title_shape_id is not None and id(shape) == title_shape_id:
                continue

            for paragraph in shape.text_frame.paragraphs:
                text = paragraph.text
                if not text.strip():
                    continue
                paragraph_index += 1
                level = (paragraph.level + 1) if paragraph.level else None
                if level:
                    builder.add(
                        "list_item",
                        text,
                        level=level,
                        page=slide_number,
                        paragraph=paragraph_index,
                    )
                else:
                    builder.add(
                        "paragraph",
                        text,
                        page=slide_number,
                        paragraph=paragraph_index,
                    )

        if len(builder.elements) == slide_start:
            builder.warn(
                f"PPTX 第 {slide_number} 页未提取到文本或表格"
                "（可能是纯图片、图表或 SmartArt 页）。"
            )

        if not notes_warned and slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text
            if notes_text and notes_text.strip():
                builder.warn("PPTX 演讲者备注未提取（仅解析幻灯片正文）。")
                notes_warned = True


def _iter_shapes(shapes: Iterable[BaseShape]) -> Iterator[BaseShape]:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _find_title_shape(shapes: list[BaseShape]) -> BaseShape | None:
    for shape in shapes:
        if not getattr(shape, "is_placeholder", False):
            continue
        try:
            placeholder_type = shape.placeholder_format.type
        except (AttributeError, ValueError):
            continue
        if placeholder_type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            return shape
    for shape in shapes:
        if shape.has_text_frame and "title" in (shape.name or "").lower():
            return shape
    return None


def _emit_table(table, builder: ElementBuilder, *, page: int, table_group: int) -> None:
    rows = list(table.rows)
    if not rows:
        return
    first_non_empty = None
    cell_rows: list[list[str]] = []
    for row in rows:
        cells = [cell.text for cell in row.cells]
        cell_rows.append(cells)
        if first_non_empty is None and any(cell.strip() for cell in cells):
            first_non_empty = len(cell_rows) - 1
    if first_non_empty is None:
        return

    for row_index, cells in enumerate(cell_rows, start=1):
        if not any(cell.strip() for cell in cells):
            continue
        element_type = (
            "table_header" if row_index - 1 == first_non_empty else "table_row"
        )
        builder.add(
            element_type,
            table_row_text(cells),
            page=page,
            row=row_index,
            table_group=table_group,
        )
