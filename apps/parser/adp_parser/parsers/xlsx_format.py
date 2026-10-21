from __future__ import annotations

import io

from openpyxl import load_workbook

from ..builder import ElementBuilder, clean_cell, table_row_text
from ..config import Limits
from ..errors import ParserError


def parse_xlsx(data: bytes, builder: ElementBuilder, limits: Limits) -> None:
    values_book = None
    formulas_book = None
    try:
        values_book = load_workbook(
            io.BytesIO(data), read_only=True, data_only=True
        )
        formulas_book = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False
        )

        any_formula = False
        formula_fallback = False
        table_group = 0

        for sheet_index, values_sheet in enumerate(values_book.worksheets):
            formulas_sheet = formulas_book.worksheets[sheet_index]
            sheet_name = values_sheet.title
            known_rows = values_sheet.max_row or 0
            known_columns = values_sheet.max_column or 0
            row_limit = limits.max_excel_rows
            column_limit = min(
                known_columns or limits.max_excel_columns,
                limits.max_excel_columns,
            )
            if known_rows > limits.max_excel_rows:
                builder.warn(
                    f"工作表 {sheet_name!r} 超过 {limits.max_excel_rows} 行上限，"
                    "后续行已截断。"
                )
            if known_columns > limits.max_excel_columns:
                builder.warn(
                    f"工作表 {sheet_name!r} 超过 {limits.max_excel_columns} 列上限，"
                    "后续列已截断。"
                )

            value_rows = values_sheet.iter_rows(
                min_row=1, max_row=row_limit, max_col=column_limit
            )
            formula_rows = formulas_sheet.iter_rows(
                min_row=1, max_row=row_limit, max_col=column_limit
            )
            started = False
            header_emitted = False
            cells_seen = 0
            cell_limit_hit = False

            for row_index, (value_row, formula_row) in enumerate(
                zip(value_rows, formula_rows), start=1
            ):
                cells: list[str] = []
                value_cells = list(value_row)
                formula_cells = list(formula_row)
                for column_index in range(column_limit):
                    value = (
                        value_cells[column_index].value
                        if column_index < len(value_cells)
                        else None
                    )
                    formula = (
                        formula_cells[column_index].value
                        if column_index < len(formula_cells)
                        else None
                    )
                    if isinstance(formula, str) and formula.startswith("="):
                        any_formula = True
                    if value is None and formula is not None:
                        value = formula
                        if isinstance(formula, str) and formula.startswith("="):
                            formula_fallback = True
                    cells.append(clean_cell(value))

                while cells and not cells[-1]:
                    cells.pop()
                if not any(cell for cell in cells):
                    continue

                if not started:
                    table_group += 1
                    started = True
                    builder.add_heading(sheet_name, 1, sheet=sheet_name)
                element_type = "table_header" if not header_emitted else "table_row"
                builder.add(
                    element_type,
                    table_row_text(cells),
                    sheet=sheet_name,
                    row=row_index,
                    table_group=table_group,
                )
                header_emitted = True
                cells_seen += len(cells)
                if cells_seen >= limits.max_excel_cells:
                    builder.warn(
                        f"工作表 {sheet_name!r} 单元格数量超过 "
                        f"{limits.max_excel_cells} 上限，后续行已截断。"
                    )
                    cell_limit_hit = True
                    break

            if not started:
                builder.warn(f"工作表 {sheet_name!r} 无数据，已跳过。")
            if cell_limit_hit:
                continue

        if any_formula:
            if formula_fallback:
                builder.warn(
                    "工作簿包含公式；部分公式无缓存值，已返回公式文本（未计算公式结果）。"
                )
            else:
                builder.warn("工作簿包含公式，已优先返回 Excel 缓存的公式计算值。")
    except ParserError:
        raise
    except Exception as exc:
        raise ParserError(
            "unreadable_xlsx",
            f"XLSX 解析失败：{type(exc).__name__}。文件可能损坏或不是有效的 XLSX。",
            422,
        ) from exc
    finally:
        if values_book is not None:
            values_book.close()
        if formulas_book is not None:
            formulas_book.close()
