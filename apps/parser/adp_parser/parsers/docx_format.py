from __future__ import annotations

import io
import re

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from ..builder import ElementBuilder, table_row_text
from ..errors import ParserError


def parse_docx(data: bytes, builder: ElementBuilder) -> None:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ParserError(
            "unreadable_docx",
            f"DOCX 解析失败：{type(exc).__name__}。文件可能损坏或不是有效的 DOCX。",
            422,
        ) from exc

    paragraph_index = 0
    table_group = 0
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph_index += 1
            paragraph = Paragraph(child, document)
            text = _paragraph_text(paragraph)
            if not text.strip():
                continue
            heading_level = _heading_level(paragraph)
            if heading_level is not None:
                builder.add_heading(
                    text, heading_level, paragraph=paragraph_index
                )
                continue
            if _is_list_paragraph(paragraph):
                builder.add(
                    "list_item",
                    text,
                    level=_list_level(paragraph),
                    paragraph=paragraph_index,
                )
            else:
                builder.add("paragraph", text, paragraph=paragraph_index)
        elif child.tag == qn("w:tbl"):
            table_group += 1
            _emit_table(Table(child, document), builder, table_group)


def _paragraph_text(paragraph: Paragraph) -> str:
    parts: list[str] = []
    for node in paragraph._p.iter():
        if node.tag == qn("w:t"):
            parts.append(node.text or "")
        elif node.tag == qn("w:tab"):
            parts.append("\t")
        elif node.tag in {qn("w:br"), qn("w:cr")}:
            parts.append("\n")
        elif node.tag == qn("w:noBreakHyphen"):
            parts.append("-")
    text = "".join(parts).replace("\r", "\n")
    return "\n".join(line.rstrip() for line in text.split("\n")).strip()


def _heading_level(paragraph: Paragraph) -> int | None:
    style_name = (paragraph.style.name or "").strip() if paragraph.style else ""
    match = re.match(r"^heading\s*([1-9])$", style_name, re.IGNORECASE)
    if match:
        return min(9, int(match.group(1)))
    match = re.match(r"^(?:标题|見出し|見出し\s*[1-9])\s*([1-9])?$", style_name)
    if match:
        return min(9, int(match.group(1) or 1))
    if style_name.lower() == "title":
        return 1

    p_pr = paragraph._p.pPr
    if p_pr is not None:
        outline = p_pr.find(qn("w:outlineLvl"))
        if outline is not None:
            value = outline.get(qn("w:val"))
            if value is not None:
                try:
                    level = int(value) + 1
                except ValueError:
                    level = 0
                if 1 <= level <= 6:
                    return level
    return None


def _is_list_paragraph(paragraph: Paragraph) -> bool:
    style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
    if "list" in style_name or "列表" in style_name:
        return True
    p_pr = paragraph._p.pPr
    return p_pr is not None and p_pr.numPr is not None


def _list_level(paragraph: Paragraph) -> int:
    p_pr = paragraph._p.pPr
    if p_pr is not None and p_pr.numPr is not None:
        ilvl = p_pr.numPr.ilvl
        if ilvl is not None and ilvl.val is not None:
            return min(9, int(ilvl.val) + 1)
    return 1


def _cell_text(cell) -> str:
    parts = [_paragraph_text(paragraph) for paragraph in cell.paragraphs]
    return " ".join(part for part in parts if part.strip())


def _emit_table(table: Table, builder: ElementBuilder, table_group: int) -> None:
    rows = list(table.rows)
    if not rows:
        return
    row_cells: list[list[str]] = []
    explicit_headers: set[int] = set()
    for index, row in enumerate(rows):
        cells = [_cell_text(cell) for cell in row.cells]
        row_cells.append(cells)
        tr_pr = row._tr.find(qn("w:trPr"))
        if tr_pr is not None and tr_pr.find(qn("w:tblHeader")) is not None:
            explicit_headers.add(index)

    first_non_empty = next(
        (index for index, cells in enumerate(row_cells) if any(c.strip() for c in cells)),
        None,
    )
    if first_non_empty is None:
        return
    header_indices = explicit_headers or {first_non_empty}

    for index, cells in enumerate(row_cells):
        if not any(cell.strip() for cell in cells):
            continue
        element_type = "table_header" if index in header_indices else "table_row"
        builder.add(
            element_type,
            table_row_text(cells),
            row=index + 1,
            table_group=table_group,
        )
