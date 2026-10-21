from __future__ import annotations

import io

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle


def make_docx() -> bytes:
    document = Document()
    document.add_heading("Chapter One", level=1)
    document.add_paragraph("First body paragraph.")
    document.add_heading("Section A", level=2)
    bullet = document.add_paragraph("Bullet item", style="List Bullet")
    bullet.paragraph_format.left_indent = Pt(18)
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Age"
    table.cell(1, 0).text = "Alice"
    table.cell(1, 1).text = "30"
    table.cell(2, 0).text = "Bob"
    table.cell(2, 1).text = "25"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pptx() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Slide Title"
    textbox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(2))
    text_frame = textbox.text_frame
    text_frame.text = "First paragraph"
    paragraph = text_frame.add_paragraph()
    paragraph.text = "Second paragraph"
    shape = slide.shapes.add_table(
        2, 2, Inches(1), Inches(4), Inches(4), Inches(1)
    )
    table = shape.table
    table.cell(0, 0).text = "Column A"
    table.cell(0, 1).text = "Column B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def make_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Age"])
    sheet.append(["Alice", 30])
    sheet.append(["Bob", 25])
    empty = workbook.create_sheet("Empty")
    empty["A1"] = None
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_pdf() -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter)
    pdf.setFont("Helvetica-Bold", 20)
    pdf.drawString(72, 720, "Large Heading")
    pdf.setFont("Helvetica", 10)
    pdf.drawString(72, 690, "This is body text on page one.")
    pdf.drawString(72, 676, "It continues on the next line.")
    pdf.showPage()
    pdf.setFont("Helvetica", 10)
    pdf.drawString(72, 720, "Page two body.")
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_pdf_table() -> bytes:
    buffer = io.BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=letter)
    table = Table(
        [["Name", "Age"], ["Alice", "30"], ["Bob", "25"]],
        colWidths=[120, 80],
    )
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 1, colors.black),
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ]
        )
    )
    document.build([table])
    return buffer.getvalue()


def make_blank_pdf() -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter)
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_encrypted_pdf() -> bytes:
    source = PdfReader(io.BytesIO(make_pdf()))
    writer = PdfWriter()
    for page in source.pages:
        writer.add_page(page)
    writer.encrypt("secret-password")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()
