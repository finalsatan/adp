"""内置表格解析（CSV、TSV、Excel）。

知识库技术方案 6.2 的“内置表格解析”一行：pandas + openpyxl，按“表头所在行”
读取，合并单元格向下填充，按设置选择工作表；每个工作表一个表格，保留原始行号
用于定位。

位置约定（X04 完成记录里也写了）：

- 每个工作表一页，``page_no`` 从 1 开始，只数被选中的工作表；同时按 Docling 的
  做法建一个 ``GroupLabel.SHEET`` 的组，组名就是工作表名。
- 表格的每个网格行对应的原始行号存在表格元素的 ``meta`` 自定义字段
  ``adp__row_numbers`` 里（1 起算，长度等于 ``table.data.num_rows``）；工作表名
  存在 ``adp__sheet_name`` 里。第三期的分块代码从这里取位置。
- 表格中间的空行不删除，所以行号映射是直接的。CSV、TSV 没有工作表，
  ``adp__sheet_name`` 为空字符串，行号是文件里的行号。
"""

from __future__ import annotations

import csv
import math
from datetime import date, datetime, time
from io import StringIO
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from docling_core.types.doc.common.meta import FloatingMeta
from docling_core.types.doc.items.table.table import TableItem
from docling_core.types.doc.labels import GroupLabel

from .. import doctools
from ..errors import ParseError
from .base import ParseContext, ParseOutput
from .builtin_text import ENGINE_VERSION, decode_text

__all__ = [
    "META_NAMESPACE",
    "ROW_NUMBERS_FIELD",
    "SHEET_NAME_FIELD",
    "parse_csv",
    "parse_excel",
]

# 表格元素 meta 里自定义字段的命名空间与字段名
META_NAMESPACE = "adp"
SHEET_NAME_FIELD = "sheet_name"
ROW_NUMBERS_FIELD = "row_numbers"

# 合并单元格在 openpyxl 里的四元组 (起始行, 起始列, 结束行, 结束列)，均 0 起算
MergeRange = tuple[int, int, int, int]


def parse_csv(ctx: ParseContext) -> ParseOutput:
    """解析 CSV、TSV：整个文件一个表格元素。"""
    text = decode_text(Path(ctx.input_path).read_bytes())
    delimiter = _csv_delimiter(ctx.ext, text)
    rows = [list(row) for row in csv.reader(StringIO(text), delimiter=delimiter)]
    header_row = _header_row(ctx.table)
    table_rows = rows[header_row - 1 :]
    row_numbers = list(range(header_row, header_row + len(table_rows)))

    document = doctools.new_document(Path(ctx.input_path).stem)
    table = doctools.add_table(document, table_rows, header_rows=1)
    _set_position_meta(table, sheet_name="", row_numbers=row_numbers)
    ctx.report(1, 1)
    return ParseOutput(document=document, pages=0, engine_version=ENGINE_VERSION)


def parse_excel(ctx: ParseContext) -> ParseOutput:
    """解析 xlsx：每个选中的工作表一页、一个表格元素。"""
    if ctx.ext != "xlsx":
        raise ParseError("UNSUPPORTED_TYPE", "请先将 Excel 文件转换为 xlsx 后再解析")
    try:
        # data_only=True：公式取计算后的缓存值，没有缓存值时为 None（取空）
        frames = pd.read_excel(ctx.input_path, sheet_name=None, header=None)
        workbook = openpyxl.load_workbook(ctx.input_path, data_only=True)
    except Exception as exc:  # noqa: BLE001 - 统一转成面向用户的错误
        raise ParseError("FILE_CORRUPTED", f"Excel 文件无法解析：{exc}") from exc

    selected = _selected_sheets(ctx.table)
    names = [name for name in workbook.sheetnames if selected is None or name in selected]
    header_row = _header_row(ctx.table)

    document = doctools.new_document(Path(ctx.input_path).stem)
    total = len(names)
    for page_no, name in enumerate(names, start=1):
        grid = _frame_grid(frames.get(name))
        merges = _merge_ranges(workbook[name])
        grid, merges = _pad_grid(grid, merges)
        _fill_down(grid, merges)

        table_rows = grid[header_row - 1 :]
        row_numbers = list(range(header_row, header_row + len(table_rows)))
        table_merges = _table_merges(merges, header_row, len(table_rows))

        doctools.add_page(document, page_no, 0, 0)
        group = document.add_group(label=GroupLabel.SHEET, name=name)
        table = doctools.add_table(
            document,
            table_rows,
            header_rows=1,
            merges=table_merges,
            page_no=page_no,
            parent=group,
        )
        _set_position_meta(table, sheet_name=name, row_numbers=row_numbers)
        ctx.report(page_no, total)

    return ParseOutput(document=document, pages=total, engine_version=ENGINE_VERSION)


def _set_position_meta(
    table: TableItem,
    *,
    sheet_name: str,
    row_numbers: list[int],
) -> None:
    """把工作表名和“表格行 → 原始行号”的映射写进表格元素的 meta。"""
    if table.meta is None:
        table.meta = FloatingMeta()
    table.meta.set_custom_field(META_NAMESPACE, SHEET_NAME_FIELD, sheet_name)
    table.meta.set_custom_field(META_NAMESPACE, ROW_NUMBERS_FIELD, list(row_numbers))


def _header_row(table: dict[str, Any] | None) -> int:
    """表头所在行，1 起算，默认 1。"""
    value = (table or {}).get("header_row")
    try:
        row = int(value) if value is not None else 1
    except (TypeError, ValueError):
        row = 1
    return max(1, row)


def _selected_sheets(table: dict[str, Any] | None) -> list[str] | None:
    """要处理的工作表名列表；None 表示全部。兼容 ``all`` 字符串。"""
    sheets = (table or {}).get("sheets")
    if sheets is None:
        return None
    if isinstance(sheets, str):
        return None if sheets in ("", "all") else [sheets]
    values = [str(name) for name in sheets]
    return values or None


def _csv_delimiter(ext: str, text: str) -> str:
    """TSV 用制表符；CSV 默认逗号，用 csv.Sniffer 判断分号等实际分隔符。"""
    if ext == "tsv":
        return "\t"
    try:
        return csv.Sniffer().sniff(text[:4096], delimiters=",;").delimiter
    except csv.Error:
        return ","


def _frame_grid(frame: pd.DataFrame | None) -> list[list[str]]:
    """pandas 读出的二维数据转成字符串网格（第 0 行是 Excel 第 1 行）。"""
    if frame is None:
        return []
    return [
        [_cell_text(value) for value in row] for row in frame.itertuples(index=False, name=None)
    ]


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if value.is_integer():
            return str(int(value))
        return repr(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _merge_ranges(sheet: Any) -> list[MergeRange]:
    """工作表里的合并区域，转成 0 起算的 (起始行, 起始列, 结束行, 结束列)。"""
    return [
        (rng.min_row - 1, rng.min_col - 1, rng.max_row - 1, rng.max_col - 1)
        for rng in sheet.merged_cells.ranges
    ]


def _pad_grid(
    grid: list[list[str]], merges: list[MergeRange]
) -> tuple[list[list[str]], list[MergeRange]]:
    """把网格补齐到覆盖所有合并区域，避免合并区域超出 pandas 读出的范围。"""
    if not merges:
        return grid, merges
    max_row = max(end_row for _, _, end_row, _ in merges)
    max_col = max(end_col for _, _, _, end_col in merges)
    while len(grid) <= max_row:
        grid.append([])
    width = max_col + 1
    for row in grid:
        if len(row) < width:
            row.extend([""] * (width - len(row)))
    return grid, merges


def _fill_down(grid: list[list[str]], merges: list[MergeRange]) -> None:
    """合并单元格向下填充：把起始单元格的值复制到下方同列的每个单元格。"""
    for start_row, start_col, end_row, _end_col in merges:
        if end_row <= start_row:
            continue
        anchor = grid[start_row][start_col] if start_col < len(grid[start_row]) else ""
        for row in range(start_row + 1, end_row + 1):
            if start_col < len(grid[row]):
                grid[row][start_col] = anchor


def _table_merges(
    merges: list[MergeRange], header_row: int, num_rows: int
) -> list[tuple[int, int, int, int]]:
    """把合并区域裁到表格范围并转成 doctools 的 (行, 列, 行数, 列数)。

    纵向合并已经向下填充成普通单元格，不再保留跨行结构；横向合并保留为跨列
    单元格，跨行又跨列时每一行保留一个跨列单元格。
    """
    start = header_row - 1
    result: list[tuple[int, int, int, int]] = []
    for start_row, start_col, end_row, end_col in merges:
        if end_row < start:
            continue
        first = max(start_row, start)
        last = min(end_row, start + num_rows - 1)
        if last < first:
            continue
        col_span = end_col - start_col + 1
        if col_span <= 1:
            continue
        for row in range(first, last + 1):
            result.append((row - start, start_col, 1, col_span))
    return result
