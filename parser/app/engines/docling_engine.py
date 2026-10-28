"""Docling 标准 PDF、图片与原生 Office/HTML/Markdown 解析引擎。

docling、markdown_doc 的导入全部放在用到的地方：主进程（接口校验、注册表）不做本地
解析，不该为它加载 docling、torch 和模型代码（X18）。
"""

from __future__ import annotations

import logging
import os
import tempfile
from importlib.metadata import version
from pathlib import Path
from typing import TYPE_CHECKING
from zipfile import BadZipFile

import pypdfium2 as pdfium
from lxml.etree import XMLSyntaxError
from PIL import Image, UnidentifiedImageError

from ..errors import ParseError
from .base import ParseContext, ParseOutput
from .mhtml import extract_mhtml, read_html, write_safe_html

if TYPE_CHECKING:  # 只用于类型标注，运行时不会导入 docling（X18）
    from docling.document_converter import DocumentConverter

logger = logging.getLogger(__name__)

ENGINE_VERSION = f"docling-{version('docling')}"
_NATIVE = {"word": {"docx"}, "ppt": {"pptx"}, "html": {"html", "htm", "shtml"}, "markdown": {"md"}}
_UNSUPPORTED = {"doc", "rtf", "odt", "ppt", "odp"}
_HTML_EXTS = {"html", "htm", "shtml", "mhtml", "mht"}
_IMAGE_DIRECT = {"png", "jpg", "jpeg", "bmp", "tiff"}
_IMAGE_CONVERT = {"webp", "gif", "jp2"}
# RapidOCR 的 PP-OCR 语言码；镜像里预置的就是这一份（Dockerfile 的 --rapidocr-backend-lang）。
_OCR_LANG = ["ch"]
_SCAN_MODES = ("auto", "force")


class DoclingEngine:
    name = "docling"

    def __init__(self) -> None:
        # 每种扫描页判断一个转换器：Docling 按管线参数缓存管线，重建会重新加载模型。
        self._converters: dict[str, DocumentConverter] = {}

    def supports(self, file_type: str, ext: str) -> bool:
        return (
            file_type == "pdf"
            or ext in _IMAGE_DIRECT | _IMAGE_CONVERT
            or ext
            in {
                "docx",
                "pptx",
                "html",
                "htm",
                "shtml",
                "mhtml",
                "mht",
                "md",
                "doc",
                "rtf",
                "odt",
                "ppt",
                "odp",
            }
        )

    def availability(self) -> tuple[bool, str | None]:
        return True, None

    def _get_converter(self, scan_mode: str) -> DocumentConverter:
        # 延迟导入：主进程只用 DoclingEngine 做接口校验，不该因此加载 torch 和模型代码（X18）
        from docling.datamodel.base_models import InputFormat
        from docling.datamodel.pipeline_options import (
            OcrMode,
            PdfPipelineOptions,
            RapidOcrOptions,
            TableFormerMode,
            TableStructureOptions,
        )
        from docling.document_converter import (
            DocumentConverter,
            ImageFormatOption,
            PdfFormatOption,
        )
        from docling.pipeline.standard_pdf_pipeline import StandardPdfPipeline

        converter = self._converters.get(scan_mode)
        if converter is None:
            options = PdfPipelineOptions(
                do_ocr=True,
                do_table_structure=True,
                table_structure_options=TableStructureOptions(mode=TableFormerMode.ACCURATE),
                # 固定用 RapidOCR（CPU，onnxruntime 后端），不交给 Docling 自动挑选引擎
                # （自动挑选在没装 Tesseract/EasyOCR 的镜像里会挑到别的引擎或直接不可用）。
                ocr_options=RapidOcrOptions(backend="onnxruntime", lang=list(_OCR_LANG)),
            )
            if scan_mode == "force":
                options.ocr_options.mode = OcrMode.FULL_PAGE
            artifacts = Path(os.environ.get("DOCLING_ARTIFACTS_PATH", "/opt/docling-models"))
            if artifacts.is_dir():
                options.artifacts_path = artifacts
            converter = DocumentConverter(
                format_options={
                    InputFormat.PDF: PdfFormatOption(
                        pipeline_cls=StandardPdfPipeline, pipeline_options=options
                    ),
                    InputFormat.IMAGE: ImageFormatOption(
                        pipeline_cls=StandardPdfPipeline, pipeline_options=options
                    ),
                }
            )
            self._converters[scan_mode] = converter
        return converter

    def warmup(self) -> None:
        """把模型读进内存（X01 的 pool 在每个子进程报到 ready 前调用）。"""
        from docling.datamodel.base_models import InputFormat

        self._get_converter("auto").initialize_pipeline(InputFormat.PDF)

    def parse(self, ctx: ParseContext) -> ParseOutput:
        if not self.supports(ctx.file_type, ctx.ext):
            raise ParseError("UNSUPPORTED_TYPE", f"Docling 引擎不支持 {ctx.file_type} 类型")
        if ctx.ext in _UNSUPPORTED:
            raise ParseError("UNSUPPORTED_TYPE", "请先转换为 docx / pptx")
        if ctx.ext not in _IMAGE_DIRECT | _IMAGE_CONVERT and ctx.file_type == "image":
            raise ParseError("UNSUPPORTED_TYPE", f"Docling 不支持 {ctx.ext} 图片")
        return _parse(ctx, self)


def _page_progress(ctx: ParseContext, fallback_pages: int):
    """Docling 的页级进度回调：只关心“每页处理完”事件。"""
    from docling.datamodel.progress import ConversionProgressKind

    def callback(event) -> None:
        if getattr(event, "kind", None) != ConversionProgressKind.PAGE_COMPLETED:
            return
        total = int(getattr(event, "total_pages", 0) or fallback_pages)
        ctx.report(int(getattr(event, "completed_pages", 0)), total)

    return callback


def parse_native(ctx: ParseContext) -> ParseOutput:
    """供 builtin 引擎转发 Markdown/HTML：复用注册表里的 Docling 引擎（X18）。

    以前每个 HTML、Markdown 作业都新建一个 ``DoclingEngine``（连带新建 ``DocumentConverter``、
    丢掉已缓存的管线）；注册表里没有时（单测直接用引擎）退回到模块级共享实例。
    """
    from .registry import get_engine

    engine = get_engine("docling")
    if engine is None:
        global _native_engine
        if _native_engine is None:
            _native_engine = DoclingEngine()
        engine = _native_engine
    return _parse(ctx, engine)


_native_engine: DoclingEngine | None = None


def _markdown_document(path: Path, name: str):
    """Markdown 文件转 DoclingDocument（延迟导入 markdown_doc，X18）。"""
    from .markdown_doc import markdown_file_to_document

    return markdown_file_to_document(path, name)


def _normalize_heading_levels(document) -> None:
    """统一标题级别（延迟导入 markdown_doc，X18）。"""
    from .markdown_doc import normalize_heading_levels

    normalize_heading_levels(document)


# 出现这些异常说明输入文件本身有问题（X17 复审）
_CORRUPTED_CAUSES = (UnidentifiedImageError, BadZipFile, pdfium.PdfiumError)
# 图片被截断或像素过大时 PIL 抛的异常（X17 复审）
_IMAGE_CAUSES = (OSError, Image.DecompressionBombError)


def _cause_chain(exc: BaseException):
    """异常的 ``__cause__`` 链；带自引用保护，避免死循环。"""
    seen: set[int] = set()
    while exc is not None and id(exc) not in seen:
        seen.add(id(exc))
        yield exc
        exc = exc.__cause__


def _failure_code(exc: BaseException, *, image: bool) -> str:
    """按异常本身和 ``__cause__`` 区分文件问题（FILE_CORRUPTED）和引擎问题（ENGINE_ERROR）。

    Docling 把管线内部的错误统一包成 ``RuntimeError("Pipeline … failed")``，格式有问题的文件
    也会变成引擎错误；反过来“没有可用管线”的 ``ConversionError`` 是配置问题。只看最外层异常
    分不清，要看 ``__cause__``（X17 复审）；分不清的按引擎问题处理。
    """
    from docling.exceptions import ConversionError

    chain = list(_cause_chain(exc))
    for error in chain:
        if isinstance(error, _CORRUPTED_CAUSES) or (image and isinstance(error, _IMAGE_CAUSES)):
            return "FILE_CORRUPTED"
    for error in chain:
        if isinstance(error, ConversionError):
            # “No pipeline could be initialized” 是格式没接管线（配置问题），不是文件坏了
            return "ENGINE_ERROR" if "could be initialized" in str(error) else "FILE_CORRUPTED"
    return "ENGINE_ERROR"


def _parse(ctx: ParseContext, engine: DoclingEngine) -> ParseOutput:
    scan_mode = ctx.engine_params.get("scan_mode", "auto")
    if scan_mode not in _SCAN_MODES:
        raise ParseError("INVALID_PARAMS", "scan_mode 必须为 auto 或 force")
    input_path = Path(ctx.input_path)
    temp_root = Path(ctx.work_dir)
    temp_root.mkdir(parents=True, exist_ok=True)
    temp = None
    ocr_pages = 0
    pages = 0
    no_text_pages: set[int] = set()
    if ctx.file_type == "pdf":
        try:
            with pdfium.PdfDocument(str(input_path)) as pdf:
                pages = len(pdf)
                if scan_mode == "force":
                    ocr_pages = pages
                else:
                    no_text_pages = {
                        index
                        for index, page in enumerate(pdf, start=1)
                        if page.get_textpage().count_chars() == 0
                    }
                    ocr_pages = len(no_text_pages)
        except Exception as exc:  # noqa: BLE001
            raise ParseError("FILE_CORRUPTED", "PDF 文件无法读取") from exc
    elif ctx.file_type == "ppt" and ctx.ext == "pptx":
        from pptx import Presentation
        from pptx.exc import PackageNotFoundError

        try:
            pages = len(Presentation(str(input_path)).slides)
        except (BadZipFile, KeyError, PackageNotFoundError, XMLSyntaxError) as exc:
            # 只把“包读不出来”当成文件问题；内存不足、依赖出错等交给外层报 ENGINE_ERROR（X17 复审）
            raise ParseError("FILE_CORRUPTED", "PPT 文件无法读取") from exc
    elif ctx.file_type == "image":
        # 多帧图片（多页 TIFF、动画 GIF/WebP）由 Docling 逐帧处理，页数按帧数算，
        # 否则绕过 parser_max_pages 上限、用量统计也偏少（X18）。
        pages = _image_page_count(input_path, ctx.ext)
        ocr_pages = pages
    try:
        source = input_path
        if ctx.ext in _IMAGE_CONVERT:
            temp = tempfile.TemporaryDirectory(dir=ctx.work_dir)
            converted = Path(temp.name) / "converted.png"
            try:
                with Image.open(input_path) as image:
                    if ctx.ext == "gif":
                        image.seek(0)
                    image.convert("RGB").save(converted, "PNG")
            except (OSError, Image.DecompressionBombError) as exc:
                # 图片被截断、像素过大：是文件本身的问题，后端不要当临时错误重试（X17 复审）
                raise ParseError("FILE_CORRUPTED") from exc
            source = converted
        elif ctx.ext in {"mhtml", "mht"}:
            temp = tempfile.TemporaryDirectory(dir=ctx.work_dir)
            source = extract_mhtml(str(input_path), str(Path(temp.name) / "body.html"))
        elif ctx.ext in {"html", "htm", "shtml"}:
            # 同一安全过滤器也作用于普通 HTML，删除脚本和远程资源引用。
            temp = tempfile.TemporaryDirectory(dir=ctx.work_dir)
            source = write_safe_html(
                read_html(input_path.read_bytes()), Path(temp.name) / "safe.html"
            )
        ctx.report(0, pages)
        if ctx.ext == "md":
            # .md 与内置 markdown、MarkItDown 用同一套标题级别修正（X13）
            document = _markdown_document(source, input_path.stem)
        else:
            result = engine._get_converter(scan_mode).convert(
                str(source), progress_callback=_page_progress(ctx, pages)
            )
            document = result.document
            if ctx.file_type == "pdf" and scan_mode == "auto":
                # pdfium 的页码和 Docling 的 page_no 都是 1 起（docling 2.142.1 的
                # get_expected_page_nos 从 1 开始），直接取并集，不会重复计数（X18）
                ocr_pages = len(
                    no_text_pages
                    | {
                        page.page_no
                        for page in result.pages
                        if any(getattr(cell, "from_ocr", False) for cell in page.cells)
                    }
                )
            if ctx.ext in _HTML_EXTS:
                # Docling 的 HTML 后端把 <h1> 当文档标题（title），和 Markdown
                # 后端一样按“1 级标题是 1 级”修正（X13）
                _normalize_heading_levels(document)
            elif ctx.file_type == "ppt":
                # Docling 的 PPTX 后端把幻灯片标题输出成没有级别的 title，改成
                # 1 级 section_header，和 MarkItDown 解析同一份 PPT 的结果一致（X16）
                _normalize_heading_levels(document)
        if not pages:
            pages = len(document.pages)
        if ctx.file_type == "ppt":
            for index, page in enumerate(document.pages.values(), start=1):
                page.page_no = index
            _add_pptx_tables(input_path, document)
        ctx.report(pages, pages)
        return ParseOutput(
            document, pages=pages, ocr_pages=ocr_pages, engine_version=ENGINE_VERSION
        )
    except ParseError:
        raise
    except Exception as exc:  # noqa: BLE001
        code = _failure_code(exc, image=ctx.file_type == "image")
        if code == "FILE_CORRUPTED":
            # 文件本身的问题：说明用错误码表的固定中文文案，异常细节只写日志（X17 复审）
            logger.warning("Docling 解析失败：%s：%s", input_path.name, exc, exc_info=True)
        else:
            # 模型缺失、OCR 运行时错误、代码缺陷等：是引擎的问题，不能说文件坏了（X17）
            logger.exception("Docling 引擎执行失败：%s", input_path.name)
        raise ParseError(code) from exc
    finally:
        if temp is not None:
            temp.cleanup()


def _image_page_count(path: Path, ext: str) -> int:
    """图片的页数：多帧图片按帧数算（X18）。

    直接交给 Docling 的格式（png/jpg/bmp/tiff）由 Docling 逐帧处理，帧数就是页数；
    ``_IMAGE_CONVERT`` 里的格式先转成单帧 PNG 再交出去，只解析第一帧。读不出来时按 1 页算，
    让后面的 Docling 报出真正的错误码。
    """
    if ext in _IMAGE_CONVERT:
        return 1
    try:
        with Image.open(path) as image:
            return max(1, int(getattr(image, "n_frames", 1)))
    except (OSError, Image.DecompressionBombError):
        return 1


def _add_pptx_tables(path: Path, document) -> None:
    """Docling 2.142.1 的 PPTX 后端忽略原生表格时，用 Office 结构保留表格。

    按页配对：Docling 已经抽到表格的页先占掉名额，只补没抽到的，不再整份重加一遍（X18）。
    """
    from pptx import Presentation

    from .. import doctools

    presentation = Presentation(str(path))
    tables = [
        (page_no, shape)
        for page_no, slide in enumerate(presentation.slides, start=1)
        for shape in slide.shapes
        if shape.has_table
    ]
    if not tables:
        return
    remaining: dict[int, int] = {}
    for table in document.tables:
        page_no = _table_page_no(table)
        if page_no is not None:
            remaining[page_no] = remaining.get(page_no, 0) + 1
    for page_no, shape in tables:
        if remaining.get(page_no, 0) > 0:
            remaining[page_no] -= 1
            continue
        rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
        doctools.add_table(
            document,
            rows,
            page_no=page_no,
            bbox=(
                float(shape.left),
                float(shape.top),
                float(shape.left + shape.width),
                float(shape.top + shape.height),
            ),
        )


def _table_page_no(table) -> int | None:
    """表格所在的页码，没有页码信息（如 Word 表格）时返回 ``None``。"""
    for prov in getattr(table, "prov", None) or []:
        page_no = getattr(prov, "page_no", None)
        if page_no is not None:
            return int(page_no)
    return None
