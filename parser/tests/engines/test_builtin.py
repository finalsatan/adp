"""内置文本与表格解析的测试（X04）。

样例文件全部现场生成（见 ``samples.py``），不往仓库提交二进制。
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from app.engines.base import ParseContext
from app.engines.builtin import BuiltinEngine
from app.engines.builtin_table import META_NAMESPACE, ROW_NUMBERS_FIELD, SHEET_NAME_FIELD
from app.engines.registry import bootstrap, get_engine
from app.errors import ParseError
from tests.engines import samples

ENGINE = BuiltinEngine()


def make_ctx(
    path: Path,
    file_type: str,
    ext: str,
    *,
    table: dict | None = None,
) -> ParseContext:
    return ParseContext(
        input_path=str(path),
        file_type=file_type,
        ext=ext,
        engine_params={},
        preprocess={},
        table=table or {},
        work_dir=str(path.parent),
        report=lambda done, total: None,
    )


def meta_of(table) -> dict:
    assert table.meta is not None
    return table.meta.get_custom_part()


# ---- TXT ----


def test_txt_gbk_paragraphs(tmp_path):
    path = samples.write_gbk_txt(tmp_path / "gbk.txt")
    output = ENGINE.parse(make_ctx(path, "txt", "txt"))

    document = output.document
    assert [item.text for item in document.texts] == samples.GBK_PARAGRAPHS
    assert document.tables == []
    assert len(document.texts) == 3
    # TXT 没有页码，也不设 prov
    assert output.pages == 0
    assert all(item.prov == [] for item in document.texts)


def test_txt_utf8_bom(tmp_path):
    path = samples.write_utf8_bom_txt(tmp_path / "bom.txt")
    output = ENGINE.parse(make_ctx(path, "txt", "txt"))

    texts = [item.text for item in output.document.texts]
    assert texts == ["带 BOM 的文本", "第二段。"]
    assert not texts[0].startswith("\ufeff")


# ---- CSV、TSV ----


def test_csv_semicolon_detected(tmp_path):
    path = samples.write_semicolon_csv(tmp_path / "data.csv")
    output = ENGINE.parse(make_ctx(path, "csv", "csv"))

    document = output.document
    assert len(document.tables) == 1
    assert output.pages == 0
    table = document.tables[0]
    assert table.data.num_rows == 3
    assert table.data.num_cols == 3

    grid = table.data.grid
    assert [cell.text for cell in grid[0]] == ["姓名", "年龄", "城市"]
    assert all(cell.column_header for cell in grid[0])
    assert [cell.text for cell in grid[1]] == ["张三", "28", "北京"]
    assert [cell.text for cell in grid[2]] == ["李四", "32", "上海"]

    meta = meta_of(table)
    assert meta[f"{META_NAMESPACE}__{SHEET_NAME_FIELD}"] == ""
    assert meta[f"{META_NAMESPACE}__{ROW_NUMBERS_FIELD}"] == [1, 2, 3]


def test_tsv(tmp_path):
    path = samples.write_tsv(tmp_path / "data.tsv")
    output = ENGINE.parse(make_ctx(path, "csv", "tsv"))

    table = output.document.tables[0]
    assert table.data.num_rows == 3
    assert table.data.num_cols == 2
    grid = table.data.grid
    assert [cell.text for cell in grid[0]] == ["名称", "数量"]
    assert [cell.text for cell in grid[2]] == ["香蕉", "20"]


# ---- Excel ----


def test_excel_multi_sheet_snapshot(tmp_path):
    path = samples.write_xlsx(tmp_path / "book.xlsx")
    output = ENGINE.parse(make_ctx(path, "excel", "xlsx"))

    document = output.document
    # 快照：没有标题，表格数等于工作表数，页数等于工作表数
    assert document.texts == []
    assert len(document.tables) == 2
    assert len(document.pages) == 2
    assert output.pages == 2
    sheet_names = [
        meta_of(table)[f"{META_NAMESPACE}__{SHEET_NAME_FIELD}"] for table in document.tables
    ]
    assert sheet_names == ["明细", "汇总"]


def test_excel_merges_header_and_blank_row(tmp_path):
    path = samples.write_xlsx(tmp_path / "book.xlsx")
    output = ENGINE.parse(make_ctx(path, "excel", "xlsx", table={"header_row": 3}))

    table = output.document.tables[0]
    assert table.data.num_rows == 5  # 表头行 + 4 行数据（含中间空行）
    assert table.data.num_cols == 3
    grid = table.data.grid

    # 表头在第 3 行，其上的行被丢弃，表格第一行就是表头
    assert [cell.text for cell in grid[0]] == ["部门", "姓名", "金额"]
    assert all(cell.column_header for cell in grid[0])

    # 纵向合并 A4:A5，值向下填充
    assert grid[1][0].text == "研发部"
    assert grid[2][0].text == "研发部"

    # 表格中间的空行不删除
    assert [cell.text for cell in grid[3]] == ["", "", ""]

    # 横向合并 A7:B7，保留为跨列单元格
    assert grid[4][0].text == "市场部"
    assert grid[4][0].col_span == 2


def test_excel_row_number_mapping(tmp_path):
    """从 DoclingDocument 的任一表格行还原工作表名和原始行号，与源文件一致。"""
    path = samples.write_xlsx(tmp_path / "book.xlsx")
    output = ENGINE.parse(make_ctx(path, "excel", "xlsx", table={"header_row": 3}))
    workbook = openpyxl.load_workbook(path, data_only=True)

    assert len(output.document.tables) == 2
    for table in output.document.tables:
        meta = meta_of(table)
        sheet_name = meta[f"{META_NAMESPACE}__{SHEET_NAME_FIELD}"]
        row_numbers = meta[f"{META_NAMESPACE}__{ROW_NUMBERS_FIELD}"]
        expected = samples.EXCEL_ROWS[sheet_name]
        assert row_numbers == expected["row_numbers"]
        assert len(row_numbers) == table.data.num_rows

        sheet = workbook[sheet_name]
        grid = table.data.grid
        for source_row, column, text in expected["checks"]:
            table_row = row_numbers.index(source_row)
            assert grid[table_row][column].text == text
            assert sheet.cell(row=source_row, column=column + 1).value is not None
        if expected["blank_row"] is not None:
            table_row = row_numbers.index(expected["blank_row"])
            assert [cell.text for cell in grid[table_row]] == ["", "", ""]


def test_excel_select_one_sheet(tmp_path):
    path = samples.write_xlsx(tmp_path / "book.xlsx")
    output = ENGINE.parse(
        make_ctx(path, "excel", "xlsx", table={"sheets": ["汇总"], "header_row": 3})
    )

    assert len(output.document.tables) == 1
    assert output.pages == 1
    table = output.document.tables[0]
    assert meta_of(table)[f"{META_NAMESPACE}__{SHEET_NAME_FIELD}"] == "汇总"
    grid = table.data.grid
    assert [cell.text for cell in grid[0]] == ["类别", "合计"]
    assert [cell.text for cell in grid[1]] == ["合计", "600"]


def test_excel_position_meta_survives_export(tmp_path):
    """doc.json 导出再加载后，工作表名和行号映射仍在（第三期按它取位置）。"""
    from docling_core.types.doc.document import DoclingDocument

    path = samples.write_xlsx(tmp_path / "book.xlsx")
    output = ENGINE.parse(make_ctx(path, "excel", "xlsx", table={"header_row": 3}))
    doc_path = tmp_path / "doc.json"
    output.document.save_as_json(doc_path)

    reloaded = DoclingDocument.load_from_json(doc_path)
    table = reloaded.tables[0]
    meta = meta_of(table)
    assert meta[f"{META_NAMESPACE}__{SHEET_NAME_FIELD}"] == "明细"
    assert meta[f"{META_NAMESPACE}__{ROW_NUMBERS_FIELD}"] == [3, 4, 5, 6, 7]
    # Docling 的 sheet 组名也保留
    assert reloaded.groups[0].label.value == "sheet"
    assert reloaded.groups[0].name == "明细"


def test_excel_non_xlsx_rejected(tmp_path):
    path = tmp_path / "old.xls"
    path.write_bytes(b"not a real xls")
    with pytest.raises(ParseError) as exc:
        ENGINE.parse(make_ctx(path, "excel", "xls"))
    assert exc.value.code == "UNSUPPORTED_TYPE"


# ---- 引擎注册与能力 ----


def test_supports_and_availability():
    assert ENGINE.availability() == (True, None)
    assert ENGINE.supports("txt", "txt")
    assert ENGINE.supports("csv", "csv")
    assert ENGINE.supports("csv", "tsv")
    assert ENGINE.supports("excel", "xlsx")
    assert ENGINE.supports("excel", "xls")
    assert not ENGINE.supports("pdf", "pdf")
    assert not ENGINE.supports("markdown", "md")
    assert not ENGINE.supports("html", "html")


def test_registry_registers_builtin_without_fake():
    bootstrap(enable_fake_engine=False)
    engine = get_engine("builtin")
    assert engine is not None
    assert engine.name == "builtin"
