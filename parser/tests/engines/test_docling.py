"""Docling 引擎回归测试；样例按需生成以避免在仓库保存二进制文件。"""

from __future__ import annotations

import threading
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from app.engines.base import ParseContext
from app.engines.docling_engine import DoclingEngine
from app.engines.mhtml import extract_mhtml


def context(path: Path, kind: str, ext: str, work: Path, **params) -> ParseContext:
    return ParseContext(
        input_path=str(path),
        file_type=kind,
        ext=ext,
        engine_params=params,
        preprocess={},
        table={},
        work_dir=str(work),
        report=lambda _done, _total: None,
    )


def write_pdf(path: Path, text: str = "ordinary text layer") -> Path:
    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.drawString(72, 750, text)
    pdf.showPage()
    pdf.save()
    return path


def write_scan_pdf(path: Path) -> Path:
    image = Image.new("RGB", (1800, 600), "white")
    draw = ImageDraw.Draw(image)
    font_candidates = (
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    )
    font = next(
        (Path(candidate) for candidate in font_candidates if Path(candidate).exists()), None
    )
    font = ImageFont.truetype(str(font), 110) if font else ImageFont.load_default(size=110)
    draw.text((70, 180), "SCANNED DOCUMENT", fill="black", font=font, stroke_width=2)
    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.drawImage(ImageReader(image), 70, 420, width=450, height=150)
    pdf.showPage()
    pdf.save()
    return path


def test_pdf_scan_mode_reports_ocr_pages(tmp_path):
    engine = DoclingEngine()
    text_pdf = write_pdf(tmp_path / "text.pdf")
    scan_pdf = write_scan_pdf(tmp_path / "scan.pdf")
    ordinary = engine.parse(context(text_pdf, "pdf", "pdf", tmp_path))
    assert ordinary.ocr_pages == 0
    forced = engine.parse(context(text_pdf, "pdf", "pdf", tmp_path, scan_mode="force"))
    assert forced.ocr_pages == forced.pages == 1
    scanned = engine.parse(context(scan_pdf, "pdf", "pdf", tmp_path))
    assert scanned.ocr_pages == 1
    assert "SCANNED" in " ".join(item.text for item in scanned.document.texts).upper()


def test_docx_preserves_heading_levels_without_pages(tmp_path):
    path = tmp_path / "headings.docx"
    doc = Document()
    doc.add_heading("Top heading", level=1)
    doc.add_heading("Nested heading", level=2)
    doc.add_paragraph("Body")
    doc.save(path)
    out = DoclingEngine().parse(context(path, "word", "docx", tmp_path))
    headings = [item for item in out.document.texts if item.label.value == "section_header"]
    assert [item.text for item in headings] == ["Top heading", "Nested heading"]
    assert [item.level for item in headings] == [1, 2]
    assert all(not item.prov for item in headings)
    assert out.pages == 0


def test_pptx_has_page_numbers_and_table(tmp_path):
    path = tmp_path / "slides.pptx"
    deck = Presentation()
    for title in ("First", "Second"):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = title
    slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1))
    deck.save(path)
    out = DoclingEngine().parse(context(path, "ppt", "pptx", tmp_path))
    assert out.pages == 2
    assert len(out.document.pages) == 2
    assert len(out.document.tables) == 1


def test_webp_image_is_converted_and_recognized(tmp_path):
    path = tmp_path / "text.webp"
    image = Image.new("RGB", (900, 220), "white")
    ImageDraw.Draw(image).text((30, 70), "WEBP OCR TEST", fill="black", stroke_width=1)
    image.save(path, "WEBP")
    out = DoclingEngine().parse(context(path, "image", "webp", tmp_path))
    assert out.pages == 1
    assert "WEBP" in " ".join(item.text for item in out.document.texts).upper()


def test_mhtml_extracts_body_and_embedded_image(tmp_path):
    archive = tmp_path / "sample.mhtml"
    archive.write_bytes(
        b"MIME-Version: 1.0\r\nContent-Type: multipart/related; boundary=abc\r\n\r\n"
        b"--abc\r\nContent-Type: text/html; charset=utf-8\r\n\r\n"
        b"<html><body><h1>Archive title</h1><img src='cid:pic'></body></html>\r\n"
        b"--abc\r\nContent-Type: image/png\r\nContent-ID: <pic>\r\n"
        b"Content-Transfer-Encoding: base64\r\n\r\naGVsbG8=\r\n--abc--\r\n"
    )
    html = extract_mhtml(str(archive), str(tmp_path / "extracted.html"))
    content = html.read_text()
    assert "Archive title" in content
    assert "data:image/png;base64,aGVsbG8=" in content


def test_html_never_fetches_remote_images(tmp_path):
    requests = []

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            requests.append(self.path)
            self.send_response(200)
            self.end_headers()

        def log_message(self, *_args):
            pass

    server = HTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        path = tmp_path / "remote.html"
        path.write_text(
            f'<h1>Local body</h1><img src="http://127.0.0.1:{server.server_port}/secret">',
            encoding="utf-8",
        )
        out = DoclingEngine().parse(context(path, "html", "html", tmp_path))
        assert any("Local body" in item.text for item in out.document.texts)
        assert requests == []
    finally:
        server.shutdown()
        thread.join(timeout=2)
        server.server_close()
