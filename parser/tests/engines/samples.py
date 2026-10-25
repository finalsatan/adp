"""X04 测试用的小样例文件，用脚本现场生成，避免往仓库里提交二进制。

所有样例都很小（几 KB），生成在 pytest 的 ``tmp_path`` 下。
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

# GBK 编码的 TXT：三个段落，第一段有两行
GBK_TEXT = "第一段第一行\n第一段第二行\n\n第二段内容。\n\n第三段内容。"
GBK_PARAGRAPHS = ["第一段第一行\n第一段第二行", "第二段内容。", "第三段内容。"]

# 分号分隔的 CSV
SEMICOLON_CSV = "姓名;年龄;城市\n张三;28;北京\n李四;32;上海\n"

# 制表符分隔的 TSV
TSV_TEXT = "名称\t数量\n苹果\t10\n香蕉\t20\n"

# Excel 里的原始行号，供测试对照
EXCEL_ROWS = {
    "明细": {
        "sheet": "明细",
        "row_numbers": [3, 4, 5, 6, 7],
        "header": ["部门", "姓名", "金额"],
        # 每个表格行取 (原始行号, 列, 期望文本)；空行用 None
        "checks": [
            (4, 1, "张三"),
            (5, 1, "李四"),
            (7, 0, "市场部"),
            (7, 2, "300"),
        ],
        "blank_row": 6,
    },
    "汇总": {
        "sheet": "汇总",
        "row_numbers": [3, 4],
        "header": ["类别", "合计"],
        "checks": [(4, 0, "合计"), (4, 1, "600")],
        "blank_row": None,
    },
}


def write_gbk_txt(path: Path) -> Path:
    path.write_bytes(GBK_TEXT.encode("gbk"))
    return path


def write_utf8_bom_txt(path: Path) -> Path:
    path.write_bytes(("带 BOM 的文本\n\n第二段。").encode("utf-8-sig"))
    return path


def write_semicolon_csv(path: Path) -> Path:
    path.write_text(SEMICOLON_CSV, encoding="utf-8")
    return path


def write_tsv(path: Path) -> Path:
    path.write_text(TSV_TEXT, encoding="utf-8")
    return path


def write_xlsx(path: Path) -> Path:
    """两个工作表，表头都在第 3 行；明细表含纵向、横向合并单元格和中间空行。"""
    workbook = openpyxl.Workbook()

    detail = workbook.active
    detail.title = "明细"
    detail["A1"] = "2024 年明细"
    detail["A3"], detail["B3"], detail["C3"] = "部门", "姓名", "金额"
    detail["A4"], detail["B4"], detail["C4"] = "研发部", "张三", 100
    detail["B5"], detail["C5"] = "李四", 200
    # 第 6 行留空
    detail["A7"], detail["C7"] = "市场部", 300
    detail.merge_cells("A4:A5")  # 纵向合并，向下填充
    detail.merge_cells("A7:B7")  # 横向合并，保留跨列

    summary = workbook.create_sheet("汇总")
    summary["A1"] = "汇总"
    summary["A3"], summary["B3"] = "类别", "合计"
    summary["A4"], summary["B4"] = "合计", 600

    workbook.save(path)
    return path
