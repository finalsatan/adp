"""X07 测试用的 Word、PPT 样例，用 python-docx / python-pptx 现场生成。

X06 的 docx、pptx 样例与本卡共用；样例还没提交前先用这里的生成器（合并时统一）。
样例都很小（几十 KB），生成在 pytest 的 ``tmp_path`` 下，不往仓库提交二进制。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

# 带多级标题的 Word 文档：三级标题、每级一个段落、一个表格
DOCX_HEADINGS = [(1, "第一章 总则"), (2, "1.1 适用范围"), (3, "1.1.1 例外")]
DOCX_PARAGRAPHS = ["本平台用于构建智能体应用。", "适用于全部知识库。", "扫描件另行处理。"]
DOCX_TABLE_HEADER = ["名称", "数量"]
DOCX_TABLE_BODY = [["苹果", "10"], ["香蕉", "20"]]

# 三页 PPT：标题、正文、表格、图片、演讲者备注
PPTX_SLIDE_WIDTH = 10  # 英寸，python-pptx 默认 4:3 版式
PPTX_SLIDE_HEIGHT = 7.5
PPTX_TITLES = {1: "平台概览", 2: "费用明细", 3: "下一步"}
PPTX_BODY = {1: "智能体开发平台\n支持知识库问答", 3: "接入更多引擎"}
PPTX_TABLE = [["项目", "金额"], ["模型", "100"], ["存储", "200"]]
PPTX_NOTES = "演讲者备注内容"


def write_docx(path: Path) -> Path:
    """多级标题 + 一个带表头的表格的 Word 文档。"""
    document = Document()
    for (level, heading), paragraph in zip(DOCX_HEADINGS, DOCX_PARAGRAPHS, strict=True):
        document.add_heading(heading, level=level)
        document.add_paragraph(paragraph)
    table = document.add_table(rows=1 + len(DOCX_TABLE_BODY), cols=2)
    table.style = "Table Grid"
    for column, text in enumerate(DOCX_TABLE_HEADER):
        table.cell(0, column).text = text
    for row, values in enumerate(DOCX_TABLE_BODY, start=1):
        for column, text in enumerate(values):
            table.cell(row, column).text = text
    document.save(path)
    return path


def write_pptx(path: Path) -> Path:
    """三页 PPT：标题加正文、标题加表格、标题加正文和图片（带备注）。"""
    presentation = Presentation()
    # 只保留标题占位符，正文用文本框，避免占位符的默认提示文字混进解析结果
    title_only = presentation.slide_layouts[5]

    first = presentation.slides.add_slide(title_only)
    first.shapes.title.text = PPTX_TITLES[1]
    _add_textbox(first, PPTX_BODY[1])

    second = presentation.slides.add_slide(title_only)
    second.shapes.title.text = PPTX_TITLES[2]
    table = second.shapes.add_table(
        rows=len(PPTX_TABLE),
        cols=2,
        left=Inches(1),
        top=Inches(2),
        width=Inches(5),
        height=Inches(2),
    ).table
    for row, values in enumerate(PPTX_TABLE):
        for column, text in enumerate(values):
            table.cell(row, column).text = text

    third = presentation.slides.add_slide(title_only)
    third.shapes.title.text = PPTX_TITLES[3]
    _add_textbox(third, PPTX_BODY[3])
    third.shapes.add_picture(
        str(_write_png(path.parent / "logo.png")),
        left=Inches(1),
        top=Inches(3.5),
        width=Inches(1),
        height=Inches(0.5),
    )
    third.notes_slide.notes_text_frame.text = PPTX_NOTES

    presentation.save(path)
    return path


def _add_textbox(slide, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1.5))
    box.text_frame.text = text


def _write_png(path: Path) -> Path:
    """一张 120x60 的白色 PNG，作为 PPT 里的图片。"""
    Image.new("RGB", (120, 60), "white").save(path, format="PNG")
    return path
