"""X06 测试用的样例文件，用脚本现场生成，避免往仓库里提交二进制。

- 文字层 PDF、扫描 PDF、webp 图片、mhtml 归档、带远程图片的 HTML 是 X06 专用。
- 多级标题 docx、带表格的 pptx 与 X07（MarkItDown）共用，两卡的断言要对得上。

正文一律用英文：reportlab 内置字体没有中文字形，OCR 模型认拉丁字母也更稳。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

# 文字层 PDF：两页，每页一行正文
TEXT_PDF_PAGES = 2
TEXT_PDF_BODY = "Ordinary text layer"

# 扫描 PDF：一页，只有图片、没有文字层
SCAN_PDF_LINES = ("SCANNED DOCUMENT", "SECOND LINE HERE", "THIRD LINE TEXT")
SCAN_PDF_WORD = SCAN_PDF_LINES[0]

# webp 图片（要先用 Pillow 转成 PNG 再解析）
# 用多行文本：实测单行图片上 Docling 的版面/OCR 会时好时坏（见卡片“问题”）。
WEBP_LINES = ("WEBP OCR TEST", "SECOND LINE HERE", "THIRD LINE TEXT")
WEBP_WORD = WEBP_LINES[0]

# 多级标题 docx：两级标题 + 一段正文，Word 没有页码（需求 4.1.8）
DOCX_HEADINGS = ["Top heading", "Nested heading"]
DOCX_BODY = "Body paragraph"

# 两页 pptx，第二页带一个 2 行 2 列的表格
PPTX_PAGE_TITLES = ["First", "Second"]
PPTX_TABLE_ROWS = 2
PPTX_TABLE_COLUMNS = 2

# mhtml 归档：正文标题 + 一张内嵌图片
MHTML_TITLE = "Archive title"
MHTML_IMAGE = b"hello"

# 带远程图片的 HTML：正文 + 一张指向本地假服务的图片
HTML_TITLE = "Local body"

_FONT_CANDIDATES = (
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
)


def _text_image(lines: tuple[str, ...], size: tuple[int, int], font_size: int) -> Image.Image:
    """生成一张白底黑字的图片，字放大一些方便 OCR 认出来（整块文字居中）。"""
    image = Image.new("RGB", size, "white")
    font_path = next((Path(path) for path in _FONT_CANDIDATES if Path(path).exists()), None)
    font = ImageFont.truetype(str(font_path), font_size) if font_path else None
    if font is None:
        font = ImageFont.load_default(size=font_size)
    draw = ImageDraw.Draw(image)
    line_height = int(font_size * 1.5)
    top = (size[1] - line_height * len(lines)) // 2
    for index, line in enumerate(lines):
        left, _, right, _ = draw.textbbox((0, 0), line, font=font, stroke_width=2)
        draw.text(
            ((size[0] - (right - left)) / 2 - left, top + index * line_height),
            line,
            fill="black",
            font=font,
            stroke_width=2,
        )
    return image


def write_text_pdf(path: Path, pages: int = TEXT_PDF_PAGES) -> Path:
    """带文字层的 PDF，默认 TEXT_PDF_PAGES 页。"""
    pdf = canvas.Canvas(str(path), pagesize=A4)
    for page in range(1, pages + 1):
        pdf.drawString(72, 750, f"{TEXT_PDF_BODY} page {page}")
        pdf.showPage()
    pdf.save()
    return path


def write_scan_pdf(path: Path) -> Path:
    """只有图片、没有文字层的单页 PDF。"""
    pdf = canvas.Canvas(str(path), pagesize=A4)
    image = _text_image(SCAN_PDF_LINES, (1800, 600), 110)
    pdf.drawImage(ImageReader(image), 70, 420, 450, 150)
    pdf.showPage()
    pdf.save()
    return path


def write_webp(path: Path) -> Path:
    _text_image(WEBP_LINES, (1400, 500), 90).save(path, "WEBP")
    return path


def write_mhtml(path: Path) -> Path:
    """正文 + 内嵌（cid:）图片的 mhtml 归档。"""
    path.write_bytes(
        b"MIME-Version: 1.0\r\nContent-Type: multipart/related; boundary=abc\r\n\r\n"
        b"--abc\r\nContent-Type: text/html; charset=utf-8\r\n\r\n"
        b"<html><body><h1>" + MHTML_TITLE.encode() + b"</h1>"
        b"<img src='cid:pic'></body></html>\r\n"
        b"--abc\r\nContent-Type: image/png\r\nContent-ID: <pic>\r\n"
        b"Content-Transfer-Encoding: base64\r\n\r\naGVsbG8=\r\n--abc--\r\n"
    )
    return path


def write_remote_image_html(path: Path, port: int) -> Path:
    """正文 + 一张指向本地假服务的远程图片。"""
    path.write_text(
        f'<h1>{HTML_TITLE}</h1><img src="http://127.0.0.1:{port}/secret">',
        encoding="utf-8",
    )
    return path


def write_headings_docx(path: Path) -> Path:
    """两级标题 + 一段正文的 docx。"""
    doc = Document()
    doc.add_heading(DOCX_HEADINGS[0], level=1)
    doc.add_heading(DOCX_HEADINGS[1], level=2)
    doc.add_paragraph(DOCX_BODY)
    doc.save(path)
    return path


def write_table_pptx(path: Path) -> Path:
    """两页 pptx，第二页带一个 2x2 表格。"""
    deck = Presentation()
    for title in PPTX_PAGE_TITLES:
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = title
    slide.shapes.add_table(
        PPTX_TABLE_ROWS, PPTX_TABLE_COLUMNS, Inches(1), Inches(2), Inches(4), Inches(1)
    )
    deck.save(path)
    return path
