"""标题级别统一的跨引擎测试（X13）。

同一份多级标题内容分别作为 ``.md``（docling、builtin）、docx（docling、
markitdown）、HTML（docling、builtin）解析，标题数和每个标题的级别要完全一致：
``#``、Word 标题 1、``<h1>`` 都是 1 级 ``section_header``。

同一份多页 pptx 用 docling 和 markitdown 解析，幻灯片标题也都是 1 级（X16）。
"""

from __future__ import annotations

from pathlib import Path

from docling_core.types.doc.labels import DocItemLabel

from app.engines.base import ParseContext, ParseOutput
from app.engines.builtin import BuiltinEngine
from app.engines.docling_engine import DoclingEngine
from app.engines.markitdown_engine import MarkitdownEngine
from tests.engines import docling_samples

HEADINGS = ["第一章 总则", "1.1 适用范围", "1.1.1 例外"]
PARAGRAPHS = ["正文一", "正文二", "正文三"]
LEVELS = [1, 2, 3]

# 与 EXPECTED 对应的期望标题：级别从 1 开始，和 Word 的标题 1～3 一致
EXPECTED = list(zip(LEVELS, HEADINGS, strict=True))

MARKDOWN = "".join(
    f"{'#' * level} {heading}\n\n{paragraph}\n\n"
    for level, heading, paragraph in zip(LEVELS, HEADINGS, PARAGRAPHS, strict=True)
)


def parse(engine, path: Path, file_type: str, ext: str) -> ParseOutput:
    return engine.parse(
        ParseContext(
            input_path=str(path),
            file_type=file_type,
            ext=ext,
            engine_params={},
            preprocess={},
            table={},
            work_dir=str(path.parent),
            report=lambda _done, _total: None,
        )
    )


def heading_levels(document) -> list[tuple[int, str]]:
    """按文档顺序列出章节标题的（级别，文字）。"""
    return [
        (item.level, item.text)
        for item in document.texts
        if item.label == DocItemLabel.SECTION_HEADER
    ]


def no_title_items(document) -> bool:
    return all(item.label != DocItemLabel.TITLE for item in document.texts)


def write_md(tmp_path: Path) -> Path:
    path = tmp_path / "headings.md"
    path.write_text(MARKDOWN, encoding="utf-8")
    return path


def write_docx(tmp_path: Path) -> Path:
    from docx import Document

    document = Document()
    for level, heading, paragraph in zip(LEVELS, HEADINGS, PARAGRAPHS, strict=True):
        document.add_heading(heading, level=level)
        document.add_paragraph(paragraph)
    path = tmp_path / "headings.docx"
    document.save(path)
    return path


def write_html(tmp_path: Path) -> Path:
    body = "".join(
        f"<h{level}>{heading}</h{level}><p>{paragraph}</p>"
        for level, heading, paragraph in zip(LEVELS, HEADINGS, PARAGRAPHS, strict=True)
    )
    path = tmp_path / "headings.html"
    path.write_text(f"<html><body>{body}</body></html>", encoding="utf-8")
    return path


def test_markdown_and_word_heading_levels_match(tmp_path):
    """同一份多级标题内容，.md 和 docx 在各自引擎下的标题数和级别一致。"""
    md = write_md(tmp_path)
    docx = write_docx(tmp_path)
    parsed = {
        "md-docling": parse(DoclingEngine(), md, "markdown", "md"),
        "md-builtin": parse(BuiltinEngine(), md, "markdown", "md"),
        "docx-docling": parse(DoclingEngine(), docx, "word", "docx"),
        "docx-markitdown": parse(MarkitdownEngine(), docx, "word", "docx"),
    }
    for name, output in parsed.items():
        assert heading_levels(output.document) == EXPECTED, name
        assert no_title_items(output.document), name
        output.document.validate_document()


def test_html_heading_levels_match_markdown(tmp_path):
    """HTML 的 <h1>～<h3> 和 Markdown 的 #～### 级别一致。"""
    html = write_html(tmp_path)
    for name, engine in (("html-docling", DoclingEngine()), ("html-builtin", BuiltinEngine())):
        output = parse(engine, html, "html", "html")
        assert heading_levels(output.document) == EXPECTED, name
        assert no_title_items(output.document), name
        output.document.validate_document()


def test_pptx_slide_titles_match_between_engines(tmp_path):
    """同一份多页 pptx：两个引擎的幻灯片标题都是 1 级，文字和页码一一对应（X16）。"""
    path = docling_samples.write_table_pptx(tmp_path / "slides.pptx")
    # 每张幻灯片的标题占位符文字和它所在的页（1 起）
    expected = [
        (1, title, page) for page, title in enumerate(docling_samples.PPTX_PAGE_TITLES, start=1)
    ]
    for name, engine in (("ppt-docling", DoclingEngine()), ("ppt-markitdown", MarkitdownEngine())):
        output = parse(engine, path, "ppt", "pptx")
        assert [
            (item.level, item.text, item.prov[0].page_no)
            for item in output.document.texts
            if item.label == DocItemLabel.SECTION_HEADER
        ] == expected, name
        assert no_title_items(output.document), name
        output.document.validate_document()


def test_html_title_tag_is_not_a_heading(tmp_path):
    """HTML 的 <title> 是文档元数据（furniture），不参与标题级别修正。"""
    path = tmp_path / "title.html"
    path.write_text(
        "<html><head><title>网页标题</title></head><body><h1>正文标题</h1></body></html>",
        encoding="utf-8",
    )
    output = parse(DoclingEngine(), path, "html", "html")
    assert heading_levels(output.document) == [(1, "正文标题")]
    # <title> 留在 furniture 内容层，由统一后处理删除（X03）
    assert [
        (item.text, item.content_layer.value)
        for item in output.document.texts
        if item.label == DocItemLabel.TITLE
    ] == [("网页标题", "furniture")]
